"""Excel and Word document creation."""

from __future__ import annotations

import asyncio
from pathlib import Path
from typing import Any, ClassVar

from pydantic import Field, field_validator

from jarvis.safety.policy import Risk
from jarvis.tools.base import Tool, ToolArgs, ToolContext, ToolError, ToolResult

Cell = str | int | float | bool | None


class SheetSpec(ToolArgs):
    name: str = Field(min_length=1, max_length=31, description="Sheet tab name.")
    headers: list[str] = Field(min_length=1)
    rows: list[list[Cell]] = Field(default_factory=list)
    column_widths: list[int] | None = Field(
        default=None, description="Optional widths in characters; auto-sized if omitted."
    )

    @field_validator("name")
    @classmethod
    def _valid_sheet_name(cls, v: str) -> str:
        if any(c in v for c in "[]:*?/\\"):
            raise ValueError("sheet names cannot contain [ ] : * ? / \\")
        return v


class CreateExcelArgs(ToolArgs):
    path: str = Field(description="Target .xlsx path.")
    sheets: list[SheetSpec] = Field(min_length=1)
    overwrite: bool = False


def _write_xlsx(path: Path, sheets: list[SheetSpec]) -> list[tuple[str, int]]:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for spec in sheets:
        ws = wb.create_sheet(spec.name)
        ws.append(spec.headers)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in spec.rows:
            ws.append(list(row))
        ws.freeze_panes = "A2"
        if spec.rows:
            ws.auto_filter.ref = ws.dimensions
        for idx in range(1, len(spec.headers) + 1):
            letter = get_column_letter(idx)
            if spec.column_widths and idx <= len(spec.column_widths):
                width = spec.column_widths[idx - 1]
            else:
                values = [spec.headers[idx - 1], *(r[idx - 1] for r in spec.rows if len(r) >= idx)]
                width = min(60, max(10, *(len(str(v)) + 2 for v in values if v is not None)))
            ws.column_dimensions[letter].width = width
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))

    # Verify by reading back.
    check = load_workbook(str(path), read_only=True)
    try:
        return [(ws.title, (ws.max_row or 1) - 1) for ws in check.worksheets]
    finally:
        check.close()


class CreateExcel(Tool[CreateExcelArgs]):
    name = "create_excel"
    description = (
        "Create a formatted Excel workbook (.xlsx) with one or more sheets: bold header row, "
        "frozen header, filters, sized columns. Use it for trackers, tables and reports."
    )
    args_model: ClassVar[type[ToolArgs]] = CreateExcelArgs
    risk = Risk.WRITE

    def risk_for(self, args: CreateExcelArgs, ctx: ToolContext) -> Risk:
        return Risk.DESTRUCTIVE if ctx.guard.resolve(args.path).exists() else Risk.WRITE

    def summarize(self, args: CreateExcelArgs) -> str:
        rows = sum(len(s.rows) for s in args.sheets)
        return f"Create Excel {args.path} ({len(args.sheets)} sheet(s), {rows} rows)"

    async def run(self, args: CreateExcelArgs, ctx: ToolContext) -> ToolResult:
        path = ctx.guard.resolve(args.path)
        if path.suffix.lower() != ".xlsx":
            raise ToolError("path must end with .xlsx")
        if path.exists() and not args.overwrite:
            raise FileExistsError(f"{path} already exists; set overwrite=true to replace it")
        for s in args.sheets:
            for i, row in enumerate(s.rows):
                if len(row) > len(s.headers):
                    raise ToolError(f"sheet {s.name!r} row {i + 1} has more cells than headers")
        counts = await asyncio.to_thread(_write_xlsx, path, args.sheets)
        summary = ", ".join(f"{name}: {n} rows" for name, n in counts)
        return ToolResult(f"Created {path} and verified it opens ({summary}).")


class DocSection(ToolArgs):
    heading: str | None = None
    level: int = Field(default=1, ge=1, le=3)
    paragraphs: list[str] = Field(default_factory=list)
    bullets: list[str] = Field(default_factory=list)
    table: list[list[str]] | None = Field(
        default=None, description="Optional table; the first row is the header."
    )


class CreateWordArgs(ToolArgs):
    path: str = Field(description="Target .docx path.")
    title: str
    sections: list[DocSection] = Field(min_length=1)
    overwrite: bool = False


def _write_docx(path: Path, args: CreateWordArgs) -> int:
    import docx

    doc = docx.Document()
    doc.add_heading(args.title, level=0)
    for sec in args.sections:
        if sec.heading:
            doc.add_heading(sec.heading, level=sec.level)
        for para in sec.paragraphs:
            doc.add_paragraph(para)
        for bullet in sec.bullets:
            doc.add_paragraph(bullet, style="List Bullet")
        if sec.table:
            cols = max(len(r) for r in sec.table)
            table = doc.add_table(rows=0, cols=cols)
            table.style = "Light Grid Accent 1"
            for r in sec.table:
                cells = table.add_row().cells
                for i, value in enumerate(r):
                    cells[i].text = value
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return len(docx.Document(str(path)).paragraphs)


class CreateWord(Tool[CreateWordArgs]):
    name = "create_word_document"
    description = (
        "Create a Word document (.docx) with a title and sections of headings, paragraphs, "
        "bullet lists and tables. Use it for summaries, reports and memos."
    )
    args_model: ClassVar[type[ToolArgs]] = CreateWordArgs
    risk = Risk.WRITE

    def risk_for(self, args: CreateWordArgs, ctx: ToolContext) -> Risk:
        return Risk.DESTRUCTIVE if ctx.guard.resolve(args.path).exists() else Risk.WRITE

    def summarize(self, args: CreateWordArgs) -> str:
        return f"Create Word document {args.path} ({len(args.sections)} sections)"

    async def run(self, args: CreateWordArgs, ctx: ToolContext) -> ToolResult:
        path = ctx.guard.resolve(args.path)
        if path.suffix.lower() != ".docx":
            raise ToolError("path must end with .docx")
        if path.exists() and not args.overwrite:
            raise FileExistsError(f"{path} already exists; set overwrite=true to replace it")
        paragraphs = await asyncio.to_thread(_write_docx, path, args)
        return ToolResult(f"Created {path} and verified it opens ({paragraphs} paragraphs).")


OFFICE_TOOLS: list[Tool[Any]] = [CreateExcel(), CreateWord()]
