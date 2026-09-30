"""Text extraction for common document formats."""

from __future__ import annotations

from pathlib import Path

TEXT_SUFFIXES = {
    ".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".html", ".htm", ".log", ".ini", ".cfg",
    ".toml", ".yaml", ".yml", ".py", ".js", ".ts", ".java", ".c", ".cpp", ".h", ".cs", ".go",
    ".rs", ".sql", ".ps1", ".bat", ".sh", ".css", ".rtf",
}  # fmt: skip


def _read_docx(path: Path) -> str:
    import docx

    doc = docx.Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _read_xlsx(path: Path, max_chars: int) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    out: list[str] = []
    size = 0
    try:
        for ws in wb.worksheets:
            out.append(f"## Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                line = " | ".join("" if v is None else str(v) for v in row).rstrip(" |")
                if line:
                    out.append(line)
                    size += len(line)
                if size > max_chars:
                    return "\n".join(out)
    finally:
        wb.close()
    return "\n".join(out)


def _read_pdf(path: Path, max_chars: int) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    out: list[str] = []
    size = 0
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        out.append(f"--- page {i} ---\n{text}")
        size += len(text)
        if size > max_chars:
            break
    return "\n".join(out)


def extract_text(path: Path, max_chars: int) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        text = _read_docx(path)
    elif suffix in {".xlsx", ".xlsm"}:
        text = _read_xlsx(path, max_chars)
    elif suffix == ".pdf":
        text = _read_pdf(path, max_chars)
    else:
        raw = path.read_bytes()[: max_chars * 4]
        if suffix not in TEXT_SUFFIXES and b"\x00" in raw[:8192]:
            raise ValueError(f"{path.name} looks like a binary file; cannot read it as text")
        text = raw.decode("utf-8", errors="replace")
    if len(text) > max_chars:
        text = text[:max_chars] + f"\n[... truncated at {max_chars} characters ...]"
    return text
